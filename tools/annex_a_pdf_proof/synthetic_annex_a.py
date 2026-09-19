"""Offline synthetic Annex A R2 layout proof; NOT a production document API.

Requires python-docx and the separately supplied original hash-locked R2 DOCX.
Only ten fixed synthetic cases are supported: five 7x7 and five Regular.
Existing SPINA schedule and projection modules are imported, never copied into
this tool or reimplemented. Run from a checkout with gilbic_backend/src on
PYTHONPATH. Conversion/visual QA are separate from this DOCX assembly; no
approval, signature or DB write occurs.
"""

from __future__ import annotations

import argparse
from datetime import date
from decimal import Decimal
import hashlib
from pathlib import Path

from docx import Document
from docx.shared import Pt

from gilbic_backend.annex_a_template_binding import (
    row_values as row_values,
    all_paragraphs,
    replace_tokens,
    set_text,
    _display_rows,
    fill_table,
)
from gilbic_backend.annex_a_schedule_projection import (
    project_regular_annex_a,
    project_seven_by_seven_annex_a,
)
from gilbic_backend.contract_schedule_engine import (
    generate_contract_installments,
    prepare_regular_contract_components,
)
from gilbic_backend.seven_by_seven_signed_schedule import (
    generate_signed_seven_by_seven_schedule,
)

TEMPLATE_SHA256 = "80ef81aec3141f9c96a5d78f1edd1e7367c8a6be9ab7dca92e81b07756217ddb"
NOTICE = "SYNTHETIC TEST COPY - NOT FOR SIGNING OR RELEASE"
CASES = {
    1: ("1000.00", "1007.00"),
    7: ("1000.00", "150.00"),
    8: ("1000.00", "140.00"),
    60: ("3000.00", "71.00"),
    104: ("3000.00", "50.00"),
}
# Fixed proof terms, not product defaults or conversions of a 120-day contract.
# Calendar exclusions are explicit approved rows below, never inferred here.
REGULAR_CASES = {
    "daily": (120, "50.00", date(2026, 9, 14), "Daily (synthetic)", "REGULAR-120"),
    "weekly": (16, "375.00", date(2026, 9, 13), "Weekly (synthetic)", "REGULAR-W16"),
    "semi_monthly": (
        8,
        "750.00",
        date(2027, 1, 15),
        "Semi-monthly (15/30; synthetic)",
        "REGULAR-SM8",
    ),
    "monthly": (4, "1500.00", date(2027, 1, 31), "Monthly (synthetic)", "REGULAR-M4"),
    "custom": (
        4,
        "1200.00",
        date(2027, 1, 29),
        "Custom approved dates (synthetic)",
        "REGULAR-C4",
    ),
}
KEYS = (
    "borrower_name",
    "company_line",
    "document_id",
    "loan_id",
    "cif_version",
    "packet_id",
    "packet_version",
    "schedule_id",
    "schedule_version",
    "generated_at",
    "approval_evidence",
)


class SyntheticProofError(ValueError):
    """Synthetic proof inputs are incomplete or unsafe to assemble."""


def make_case(count: int) -> dict:
    if type(count) is not int or count not in CASES:
        raise SyntheticProofError(
            "Only the five fixed 7x7 synthetic cases are supported."
        )
    principal, daily = CASES[count]
    source = generate_signed_seven_by_seven_schedule(
        original_principal=principal,
        agreed_daily_payment=daily,
        daily_interest_per_1000="7.00",
        first_due_date=date(2026, 9, 13),
    )
    projection = project_seven_by_seven_annex_a(
        original_principal=Decimal(principal),
        installments=source,
        expected_installment_count=count,
        expected_first_due_date=date(2026, 9, 13),
        expected_maturity_date=source[-1].due_date,
        expected_total_payable=sum(
            (r.contractual_amount for r in source), Decimal("0.00")
        ),
    )
    return {
        "count": count,
        "principal": principal,
        "daily": daily,
        "source": source,
        "projection": projection,
    }


def make_regular_case(*, payment_frequency: str = "daily") -> dict:
    """Select one fixed Regular proof fixture; never accept arbitrary loan terms."""

    if payment_frequency not in REGULAR_CASES:
        raise SyntheticProofError(
            "Only the five fixed Regular synthetic cases are supported."
        )

    principal = "5000.00"
    interest = "1000.00"
    total = "6000.00"
    count, daily, first_due_date, _, _ = REGULAR_CASES[payment_frequency]
    if payment_frequency == "custom":
        # Already-approved synthetic dates and amounts; NOT a holiday calendar.
        # Jan 30/31, Feb 1/7/8 are excluded in this example before signing.
        generic = generate_contract_installments(
            payment_frequency="custom",
            contractual_total=Decimal(total),
            custom_installments=(
                (date(2027, 1, 29), Decimal("1200.00")),
                (date(2027, 2, 2), Decimal("900.00")),
                (date(2027, 2, 5), Decimal("1800.00")),
                (date(2027, 2, 9), Decimal("2100.00")),
            ),
        )
    else:
        generic = generate_contract_installments(
            payment_frequency=payment_frequency,
            contractual_total=Decimal(total),
            first_due_date=first_due_date,
            installment_count=count,
            regular_installment_amount=Decimal(daily),
        )
    source = prepare_regular_contract_components(
        installments=generic,
        original_principal=Decimal(principal),
        contractual_interest=Decimal(interest),
    )
    projection = project_regular_annex_a(
        original_principal=Decimal(principal),
        installments=source,
        expected_installment_count=count,
        expected_first_due_date=first_due_date,
        expected_maturity_date=source[-1].due_date,
        expected_total_payable=Decimal(total),
    )
    return {
        "product": "Regular",
        "payment_frequency": payment_frequency,
        "count": count,
        "principal": principal,
        "interest": interest,
        "rate_percent": f"{(Decimal(interest) / Decimal(principal) * Decimal('100')):.2f}",
        "daily": daily,
        "source": source,
        "projection": projection,
    }


def regular_specimen_stem(payment_frequency: str) -> str:
    count = REGULAR_CASES[payment_frequency][0]
    if payment_frequency == "daily":
        return "SPINA_Annex_A_SYNTHETIC_REGULAR_120_rows"
    return (
        f"SPINA_Annex_A_SYNTHETIC_REGULAR_{payment_frequency.upper()}_{count:03d}_rows"
    )


def regular_specimen_identity(payment_frequency: str) -> str:
    # Compact references preserve the existing footer geometry.
    return REGULAR_CASES[payment_frequency][4]


def synthetic_context(count: int | str) -> dict[str, str]:
    return {
        "borrower_name": f"SYNTHETIC TEST BORROWER {count}",
        "company_line": "SYNTHETIC OFFICE - NOT A REGISTERED ADDRESS",
        "document_id": f"SYN-ANNEX-{count}",
        "loan_id": f"SYN-LOAN-{count}",
        "cif_version": "SYN-CIF-v1",
        "packet_id": f"SYN-PACKET-{count}",
        "packet_version": "SYN-v1",
        "schedule_id": f"SYN-SCHEDULE-{count}",
        "schedule_version": "SYN-v1",
        "generated_at": "2026-09-12 21:35 +08:00 (synthetic)",
        "approval_evidence": f"SYN-APPROVAL-{count} / NOT APPROVED OR SIGNED",
    }


def _expected_case(case: dict) -> dict:
    product = case.get("product", "7x7")
    if product == "Regular":
        return make_regular_case(
            payment_frequency=case.get("payment_frequency", "daily")
        )
    if product == "7x7":
        return make_case(case["count"])
    raise SyntheticProofError("Unsupported synthetic product fixture.")


def _display_terms(case: dict, result) -> dict[str, str]:
    if case.get("product") == "Regular":
        frequency = case.get("payment_frequency", "daily")
        return {
            "product": "Regular Cash Loan (synthetic only)",
            "rate": f"{case['rate_percent']}% fixed contractual interest (synthetic fixture)",
            "basis": "Fixed contractual interest on original principal; synthetic fixture only",
            "frequency": REGULAR_CASES[frequency][3],
            "term": f"{case['count']} scheduled {frequency.replace('_', '-')} installments (synthetic)",
            "allocation": (
                "Synthetic Regular principal/interest schedule rows only; "
                "no payment-allocation rule represented"
            ),
        }
    count = case["count"]
    return {
        "product": "7x7 Cash Loan (synthetic only)",
        "rate": f"PHP {result.rows[0].interest_component:,.2f} per day (synthetic)",
        "basis": "Fixed on original principal; synthetic fixture only",
        "frequency": "Daily (synthetic)",
        "term": f"{count} scheduled daily installment{'s' if count != 1 else ''} (synthetic)",
        "allocation": (
            "Past Due Interest; Today's Interest; Past Due Principal; "
            "Today's Principal; Advance (synthetic disclosure)"
        ),
    }


def build_docx(template: Path, output: Path, case: dict, context: dict) -> None:
    template, output = Path(template), Path(output)
    if template.resolve() == output.resolve() or output.exists():
        raise SyntheticProofError(
            "Never overwrite the source template or an existing output."
        )
    data = template.read_bytes()
    if hashlib.sha256(data).hexdigest() != TEMPLATE_SHA256:
        raise SyntheticProofError("The exact approved Annex A R2 template is required.")
    for key in KEYS:
        value = context.get(key)
        if (
            not isinstance(value, str)
            or not value.strip()
            or "{" in value
            or "}" in value
        ):
            raise SyntheticProofError(f"Missing or unresolved synthetic context: {key}")
    if not context["borrower_name"].startswith("SYNTHETIC "):
        raise SyntheticProofError("This tool accepts synthetic specimens only.")
    count = case["count"]
    # Rebuild only a fixed test fixture; never accept a modified case as a loan.
    expected = _expected_case(case)
    if case != expected:
        raise SyntheticProofError("Modified synthetic fixture is not supported.")
    result = case["projection"]
    terms = _display_terms(case, result)
    display_rows = _display_rows(case)
    regular = case.get("product") == "Regular"
    doc = Document(template)
    tables = doc.tables
    if len(tables) != 6 or [len(t.columns) for t in tables] != [2, 7, 7, 2, 2, 2]:
        raise SyntheticProofError("Unexpected approved-template structure.")
    fill_table(tables[1], display_rows[:7])
    if count > 7:
        fill_table(tables[2], display_rows[7:])
    else:
        tables[2]._element.getparent().remove(tables[2]._element)
    # Keep totals and the borrower name/signature/approval table together.
    for table in (tables[3], tables[5]):
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    mapping = {
        "{Controlled SEC company identity and registered office address}": context[
            "company_line"
        ],
        "{Borrower Full Name from Locked Loan}": context["borrower_name"],
        "{Borrower Full Name}": context["borrower_name"],
        "{Loan Account No.}": context["loan_id"],
        "{Locked CIF Version}": context["cif_version"],
        "{Packet ID}": context["packet_id"],
        "{Packet ID and Version}": context["packet_id"]
        + " / "
        + context["packet_version"],
        "{Schedule ID and Version}": context["schedule_id"]
        + " / "
        + context["schedule_version"],
        "{Regular Cash Loan or 7x7 Cash Loan}": terms["product"],
        "{Approved Gross Principal}": f"{Decimal(case['principal']):,.2f}",
        "{Exact Approved Rate and Period / Daily Amount}": terms["rate"],
        "{Exact Approved Interest Basis and Method}": terms["basis"],
        "{Exact Approved Payment Frequency}": terms["frequency"],
        "{Exact Approved Contractual Term}": terms["term"],
        "{Agreed Installment Amount}": f"{Decimal(case['daily']):,.2f}",
        "{N - Complete Approved Installment Count}": str(count),
        "{First Contractual Due Date}": result.rows[0].due_date.isoformat(),
        "{Last Contractual Installment Date}": result.maturity_date.isoformat(),
        "{Total under Complete Contractual Schedule}": f"{result.total_due:,.2f}",
        "{Total Principal Components}": f"{result.total_principal:,.2f}",
        "{Total Interest Components}": f"{result.total_interest:,.2f}",
        "{Total Scheduled Charges / 0.00}": "0.00",
        "{Sum of All Contractual Installments}": f"{result.total_due:,.2f}",
        "{Final Scheduled Remaining Principal}": display_rows[-1][-1],
        "{Exact Product-Specific Allocation from Locked Disclosure}": terms[
            "allocation"
        ],
        "{Exact Applicable Locked Disclosure Rule / None}": "None / PHP 0.00 (synthetic disclosure; no penalty enabled)",
        "{Authenticated Approver / Locked Approval Reference}": context[
            "approval_evidence"
        ],
        "{Document ID}": context["document_id"],
        "{Template Version}": "Annex A R2",
        "{Packet Version}": context["packet_version"],
        "{Schedule Version}": context["schedule_version"],
        "{Server Date/Time and Zone}": context["generated_at"],
    }
    if regular and case.get("payment_frequency") == "custom":
        # Do not present the first custom amount as a fixed installment amount.
        mapping = {
            "PHP {Agreed Installment Amount}": "Variable; see complete schedule",
            **mapping,
        }
    for p in list(doc.paragraphs):
        text = p.text
        if text.startswith("Working layout R2"):
            set_text(
                p,
                NOTICE
                + ". Test data only; no real loan, approval, signature or cash receipt is represented.",
            )
        elif text.startswith("Template instruction:"):
            p._element.getparent().remove(p._element)
        elif text == "II. CONTRACTUAL INSTALLMENTS - FIRST-PAGE LAYOUT":
            set_text(p, "II. CONTRACTUAL INSTALLMENTS")
        elif text.startswith("All monetary values are in Philippine pesos"):
            set_text(
                p,
                "All monetary values are in Philippine pesos (PHP). Every installment is included, with continuation rows where needed.",
            )
        elif text.startswith("II. CONTRACTUAL INSTALLMENTS - CONTINUATION"):
            if count <= 7:
                p._element.getparent().remove(p._element)
            else:
                set_text(p, "II. CONTRACTUAL INSTALLMENTS - CONTINUATION")
        elif text.startswith("Loan: {Loan Account No.}") and count <= 7:
            p._element.getparent().remove(p._element)
        elif text.startswith("Row N is the actual"):
            set_text(p, text.replace("Row N", f"Row {count}"))
        elif text == "IV. BORROWER ACKNOWLEDGMENT" or text.startswith(
            "I acknowledge receipt and review of the complete Schedule"
        ):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
        elif count <= 7 and p._p.xpath('.//w:br[@w:type="page"]'):
            p._element.getparent().remove(p._element)
    for p in all_paragraphs(doc):
        replace_tokens(p, mapping)
        if regular:
            if p.text.startswith("* Scheduled Remaining Principal:"):
                set_text(
                    p,
                    (
                        "* Remaining Total Payable: Includes scheduled principal and interest "
                        "remaining after this installment, assuming all scheduled payments are "
                        "made fully and on time. This is not a live account balance or payoff "
                        "quote and is not proof of payment."
                    ),
                )
            else:
                replace_tokens(
                    p,
                    {
                        "Scheduled Remaining Principal*": "Remaining Total Payable*",
                        "FINAL SCHEDULED REMAINING PRINCIPAL": "FINAL REMAINING TOTAL PAYABLE",
                    },
                )
        if p.text == "____________________________ / ______________":
            set_text(p, "UNSIGNED TEST ONLY - DO NOT SIGN")
    # Use the existing last blank header paragraph, leaving logo bytes untouched.
    p = doc.sections[0].header.paragraphs[-1]
    set_text(p, NOTICE)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        run.font.size = Pt(7)
        run.bold = True
    if any("{" in p.text or "}" in p.text for p in all_paragraphs(doc)):
        raise SyntheticProofError(
            "Unresolved template fields remain; no output written."
        )
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    for count in CASES:
        path = args.output_dir / f"SPINA_Annex_A_SYNTHETIC_{count:03d}_rows.docx"
        build_docx(args.template, path, make_case(count), synthetic_context(count))
        print(path)
    for frequency in REGULAR_CASES:
        regular = make_regular_case(payment_frequency=frequency)
        path = args.output_dir / (regular_specimen_stem(frequency) + ".docx")
        context = synthetic_context(regular_specimen_identity(frequency))
        build_docx(args.template, path, regular, context)
        print(path)


if __name__ == "__main__":
    main()
