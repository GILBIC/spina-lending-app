"""Explicit local proof tests; require the original hash-locked Annex A R2.

Run with SPINA_ANNEX_TEMPLATE set; missing assets fail rather than skip.
These tests are not an application integration or production authorization.
"""
from copy import deepcopy
from datetime import date, timedelta
from decimal import Decimal
from importlib import import_module
from importlib.util import find_spec
import os
from pathlib import Path
import sys
import zipfile

from docx import Document
import pytest

KEYS = (
    'borrower_name', 'company_line', 'document_id', 'loan_id', 'cif_version',
    'packet_id', 'packet_version', 'schedule_id', 'schedule_version',
    'generated_at', 'approval_evidence',
)


def module():
    assert find_spec('synthetic_annex_a') is not None, 'Synthetic template binder is missing (TDD RED)'
    return import_module('synthetic_annex_a')


def template():
    p = Path(os.environ['SPINA_ANNEX_TEMPLATE'])
    assert p.is_file(), 'The approved original Annex A R2 is required'
    return p


@pytest.mark.parametrize('count', [1, 7, 8, 60, 104])
def test_every_row_exact_once_with_matching_source_money(count, tmp_path):
    m = module()
    before = template().read_bytes()
    case = m.make_case(count)
    source = case['source']
    out = tmp_path / 'proof.docx'
    m.build_docx(template(), out, case, m.synthetic_context(count))
    doc = Document(out)
    tables = [t for t in doc.tables if len(t.columns) == 7]
    actual = [[c.text for c in r.cells] for t in tables for r in t.rows[1:]]

    assert actual == [m.row_values(row) for row in case['projection'].rows]
    assert len(actual) == count
    assert actual[-1][-1] == '0.00'
    assert len(tables) == (1 if count <= 7 else 2)
    assert case['source'] == source
    assert template().read_bytes() == before


@pytest.mark.parametrize('key', KEYS)
def test_missing_context_blocks_output(key, tmp_path):
    m = module()
    context = m.synthetic_context(8)
    context.pop(key)
    out = tmp_path / 'must-not-exist.docx'
    with pytest.raises(m.SyntheticProofError):
        m.build_docx(template(), out, m.make_case(8), context)
    assert not out.exists()


@pytest.mark.parametrize('value', ['', '   ', '{unresolved}', None])
def test_blank_or_unresolved_context_blocks_output(value, tmp_path):
    m = module()
    context = m.synthetic_context(1)
    context['borrower_name'] = value
    out = tmp_path / 'must-not-exist.docx'
    with pytest.raises(m.SyntheticProofError):
        m.build_docx(template(), out, m.make_case(1), context)
    assert not out.exists()


def test_modified_template_is_rejected(tmp_path):
    m = module()
    changed = tmp_path / 'wrong.docx'
    changed.write_bytes(template().read_bytes() + b'changed')
    with pytest.raises(m.SyntheticProofError):
        m.build_docx(changed, tmp_path / 'no.docx', m.make_case(1), m.synthetic_context(1))


def test_no_placeholder_and_logo_and_word_fields_preserved(tmp_path):
    m = module()
    out = tmp_path / 'proof.docx'
    m.build_docx(template(), out, m.make_case(104), m.synthetic_context(104))
    d = Document(out)
    all_text = '\n'.join(p.text for p in m.all_paragraphs(d))
    assert '{' not in all_text and '}' not in all_text
    assert 'Template instruction:' not in all_text
    assert 'Working layout R2' not in all_text
    assert 'SYNTHETIC TEST COPY - NOT FOR SIGNING OR RELEASE' in all_text
    assert 'UNSIGNED TEST ONLY' in all_text
    assert 'Scheduled Remaining Principal' in all_text
    with zipfile.ZipFile(template()) as old, zipfile.ZipFile(out) as new:
        images = [n for n in old.namelist() if n.startswith('word/media/')]
        assert images
        assert all(old.read(n) == new.read(n) for n in images)
        footer = new.read('word/footer1.xml').decode()
        assert 'NUMPAGES' in footer and 'PAGE' in footer
        assert not any(n.startswith('word/fonts/') for n in new.namelist())


def test_output_cannot_overwrite_template(tmp_path):
    m = module()
    with pytest.raises(m.SyntheticProofError):
        m.build_docx(template(), template(), m.make_case(1), m.synthetic_context(1))


def test_totals_summary_stays_together(tmp_path):
    m = module()
    out = tmp_path / 'proof.docx'
    m.build_docx(template(), out, m.make_case(104), m.synthetic_context(104))
    d = Document(out)
    totals = next(t for t in d.tables if t.cell(0, 0).text == 'TOTAL SCHEDULED PRINCIPAL')
    for row in totals.rows[:-1]:
        assert all(p.paragraph_format.keep_with_next for cell in row.cells for p in cell.paragraphs)


def test_regular_fixture_uses_same_template_and_authoritative_component_rows(tmp_path):
    m = module()
    assert hasattr(m, 'make_regular_case'), (
        'Regular synthetic Annex A fixture is intentionally missing at this TDD RED step.'
    )
    case = m.make_regular_case()
    source_before = tuple(case['source'])
    projection_before = case['projection']
    template_before = template().read_bytes()

    assert case['product'] == 'Regular'
    assert case['count'] == 120
    assert case['projection'].total_principal == Decimal('5000.00')
    assert case['projection'].total_interest == Decimal('1000.00')
    assert case['projection'].total_due == Decimal('6000.00')

    out = tmp_path / 'regular-proof.docx'
    m.build_docx(template(), out, case, m.synthetic_context('REGULAR-120'))
    doc = Document(out)
    tables = [t for t in doc.tables if len(t.columns) == 7]
    actual = [[c.text for c in r.cells] for t in tables for r in t.rows[1:]]

    # Presentation feedback must not turn the approved components principal-first.
    assert case['source'] == source_before
    assert case['projection'] == projection_before
    assert template().read_bytes() == template_before
    assert source_before[0].principal_component == Decimal('41.67')
    assert source_before[0].interest_component == Decimal('8.33')
    assert projection_before.rows[0].scheduled_remaining_principal == Decimal('4958.33')
    assert projection_before.rows[1].scheduled_remaining_principal == Decimal('4916.67')
    assert projection_before.rows[99].scheduled_remaining_principal == Decimal('833.33')
    assert [values[:6] for values in actual] == [
        [str(row.installment_number), row.due_date.isoformat(),
         f'{row.principal_component:,.2f}', f'{row.interest_component:,.2f}',
         '0.00', f'{row.contractual_amount:,.2f}']
        for row in source_before
    ]
    assert len(actual) == 120
    # Independent expected values for this fixed 120 x PHP50 specimen only.
    expected_display = [
        f'{Decimal("6000.00") - Decimal("50.00") * number:,.2f}'
        for number in range(1, 121)
    ]
    assert [values[-1] for values in actual] == expected_display
    assert [values[-1] for values in actual[:3]] == ['5,950.00', '5,900.00', '5,850.00']
    assert actual[98][-1] == '1,050.00'
    assert actual[99][-1] == '1,000.00'
    assert actual[118][-1] == '50.00'
    assert actual[119][-1] == '0.00'
    assert all(Decimal(values[-1].replace(',', '')) > 0 for values in actual[:-1])
    # Each displayed total equals the principal and interest still scheduled.
    for number, values in enumerate(actual, start=1):
        future = source_before[number:]
        remaining_components = sum(
            (row.principal_component + row.interest_component for row in future),
            Decimal('0.00'),
        )
        assert Decimal(values[-1].replace(',', '')) == remaining_components
    all_text = '\n'.join(p.text for p in m.all_paragraphs(doc))
    assert 'Regular Cash Loan (synthetic only)' in all_text
    assert '7x7 Cash Loan (synthetic only)' not in all_text


def test_regular_remaining_total_payable_has_consistent_labels_and_schedule_note(tmp_path):
    m = module()
    out = tmp_path / 'regular-presentation-label.docx'
    m.build_docx(template(), out, m.make_regular_case(), m.synthetic_context('REGULAR-120'))
    doc = Document(out)
    tables = [t for t in doc.tables if len(t.columns) == 7]

    assert len(tables) == 2
    for table in tables:
        label = ' '.join(table.cell(0, 6).text.split())
        assert label == 'Remaining Total Payable*'
    text = ' '.join(' '.join(p.text for p in m.all_paragraphs(doc)).split()).lower()
    assert 'capital recovery balance' not in text
    assert 'includes scheduled principal and interest' in text
    assert 'assuming all scheduled payments are made fully and on time' in text
    assert 'not a live account balance or payoff quote' in text
    assert 'not proof of payment' in text
    totals = next(t for t in doc.tables if t.cell(0, 0).text == 'TOTAL SCHEDULED PRINCIPAL')
    assert [row.cells[0].text for row in totals.rows] == [
        'TOTAL SCHEDULED PRINCIPAL', 'TOTAL CONTRACTUAL INTEREST',
        'TOTAL OTHER LAWFUL SCHEDULED CHARGES', 'TOTAL AMOUNT PAYABLE',
        'FINAL REMAINING TOTAL PAYABLE',
    ]
    assert [row.cells[1].text for row in totals.rows] == [
        'PHP 5,000.00', 'PHP 1,000.00', 'PHP 0.00', 'PHP 6,000.00', 'PHP 0.00',
    ]


def test_cli_generates_five_7x7_and_five_regular_specimens(tmp_path, monkeypatch):
    m = module()
    output_dir = tmp_path / 'specimens'
    monkeypatch.setattr(
        sys,
        'argv',
        [
            'synthetic_annex_a.py',
            '--template',
            str(template()),
            '--output-dir',
            str(output_dir),
        ],
    )

    m.main()

    outputs = sorted(path.name for path in output_dir.glob('*.docx'))
    expected = {f'SPINA_Annex_A_SYNTHETIC_{count:03d}_rows.docx' for count in (1, 7, 8, 60, 104)}
    expected.add('SPINA_Annex_A_SYNTHETIC_REGULAR_120_rows.docx')
    expected.update({
        'SPINA_Annex_A_SYNTHETIC_REGULAR_WEEKLY_016_rows.docx',
        'SPINA_Annex_A_SYNTHETIC_REGULAR_SEMI_MONTHLY_008_rows.docx',
        'SPINA_Annex_A_SYNTHETIC_REGULAR_MONTHLY_004_rows.docx',
        'SPINA_Annex_A_SYNTHETIC_REGULAR_CUSTOM_004_rows.docx',
    })
    assert set(outputs) == expected


@pytest.mark.parametrize('case_id', [1, 7, 8, 60, 104, 'REGULAR-120'])
def test_acknowledgment_text_name_and_signature_stay_together(case_id, tmp_path):
    # CI specimen 8c5a9f4 split Regular printed name (page 5) from signature (6).
    m = module()
    case = m.make_regular_case() if case_id == 'REGULAR-120' else m.make_case(case_id)
    before = template().read_bytes()
    out = tmp_path / 'acknowledgment-proof.docx'
    m.build_docx(template(), out, case, m.synthetic_context(case_id))
    doc = Document(out)
    heading = next(p for p in doc.paragraphs if p.text == 'IV. BORROWER ACKNOWLEDGMENT')
    acknowledgment = next(
        p for p in doc.paragraphs
        if p.text.startswith('I acknowledge receipt and review of the complete Schedule')
    )
    signatures = next(
        t for t in doc.tables if t.cell(0, 0).text == 'Borrower Printed Name'
    )
    assert [row.cells[0].text for row in signatures.rows] == [
        'Borrower Printed Name', 'Borrower Signature / Date', 'Management Approval / Evidence',
    ]
    assert 'UNSIGNED TEST ONLY - DO NOT SIGN' in signatures.cell(1, 1).text
    for paragraph in (heading, acknowledgment):
        assert paragraph.paragraph_format.keep_with_next is True, (
            'Acknowledgment heading/text must stay with the signature table.'
        )
        assert paragraph.paragraph_format.keep_together is True
    for row in signatures.rows[:-1]:
        assert all(
            p.paragraph_format.keep_with_next is True
            for cell in row.cells for p in cell.paragraphs
        ), 'Borrower printed name, signature and approval must stay on one page.'
    assert template().read_bytes() == before


# Fixed synthetic terms, not conversions of a 120-day contract or public holidays.
# These cases exercise the existing signed-row boundary, not a holiday provider.
REGULAR_CALENDAR_CASES = (
    ('weekly', 'Weekly (synthetic)',
     tuple(date(2026, 9, 13) + timedelta(days=7 * n) for n in range(16)),
     ('375.00',) * 16),
    ('semi_monthly', 'Semi-monthly (15/30; synthetic)',
     tuple(date.fromisoformat(value) for value in (
         '2027-01-15', '2027-01-30', '2027-02-15', '2027-02-28',
         '2027-03-15', '2027-03-30', '2027-04-15', '2027-04-30',
     )), ('750.00',) * 8),
    ('monthly', 'Monthly (synthetic)',
     tuple(date.fromisoformat(value) for value in (
         '2027-01-31', '2027-02-28', '2027-03-31', '2027-04-30',
     )), ('1500.00',) * 4),
    ('custom', 'Custom approved dates (synthetic)',
     tuple(date.fromisoformat(value) for value in (
         '2027-01-29', '2027-02-02', '2027-02-05', '2027-02-09',
     )), ('1200.00', '900.00', '1800.00', '2100.00')),
)


@pytest.mark.parametrize('frequency,label,dates,amounts', REGULAR_CALENDAR_CASES)
def test_regular_calendar_specimen_preserves_exact_signed_rows(
    frequency, label, dates, amounts, tmp_path,
):
    m = module()
    # Extend the fixed fixture factory only; arbitrary live loan inputs stay rejected.
    case = m.make_regular_case(payment_frequency=frequency)
    before = deepcopy(case)
    template_before = template().read_bytes()
    source = tuple(case['source'])
    assert case['product'] == 'Regular'
    assert [r.due_date for r in source] == list(dates)
    assert [r.contractual_amount for r in source] == [Decimal(a) for a in amounts]
    assert case['count'] == len(dates)
    assert sum((r.principal_component for r in source), Decimal('0.00')) == Decimal('5000.00')
    assert sum((r.interest_component for r in source), Decimal('0.00')) == Decimal('1000.00')

    # Sundays and explicitly synthetic holidays can remain approved due dates.
    if frequency in {'weekly', 'semi_monthly', 'monthly'}:
        assert any(d.weekday() == 6 for d in dates)
        synthetic_holidays = {dates[-1]}
        assert synthetic_holidays <= {r.due_date for r in source}
    else:
        # Sunday + holiday overlap and consecutive exclusions, including the tail.
        # These exact dates were selected before signing, not moved by the renderer.
        excluded = {date.fromisoformat(value) for value in (
            '2027-01-30', '2027-01-31', '2027-02-01', '2027-02-07', '2027-02-08',
        )}
        assert date(2027, 1, 31).weekday() == 6
        assert not excluded.intersection(r.due_date for r in source)
        assert source[-1].due_date == date(2027, 2, 9)

    out = tmp_path / f'regular-{frequency}.docx'
    m.build_docx(template(), out, case, m.synthetic_context(f'REGULAR-{frequency}'))
    doc = Document(out)
    tables = [t for t in doc.tables if len(t.columns) == 7]
    actual = [[c.text for c in row.cells] for t in tables for row in t.rows[1:]]
    assert len(tables) == (1 if len(source) <= 7 else 2)
    assert len(actual) == len(source)
    for index, (row, cells) in enumerate(zip(source, actual, strict=True)):
        assert cells[:6] == [
            str(index + 1), dates[index].isoformat(),
            f'{row.principal_component:,.2f}', f'{row.interest_component:,.2f}',
            '0.00', f'{Decimal(amounts[index]):,.2f}',
        ]
        future_components = sum(
            (r.principal_component + r.interest_component for r in source[index + 1:]),
            Decimal('0.00'),
        )
        assert cells[-1] == f'{future_components:,.2f}'
    assert all(Decimal(cells[-1].replace(',', '')) > 0 for cells in actual[:-1])
    assert actual[-1][-1] == '0.00'
    for table in tables:
        assert ' '.join(table.cell(0, 6).text.split()) == 'Remaining Total Payable*'
    details = {r.cells[0].text: r.cells[1].text for r in doc.tables[0].rows}
    assert details['Payment Frequency'] == label
    assert 'daily' not in details['Contractual Loan Term'].lower()
    assert details['Number of Installments'] == str(len(dates))
    assert details['First Payment Date'] == dates[0].isoformat()
    assert details['Contractual Maturity'] == dates[-1].isoformat()
    assert details['Total Amount Payable'] == 'PHP 6,000.00'
    summary = next(t for t in doc.tables if t.cell(0, 0).text == 'TOTAL SCHEDULED PRINCIPAL')
    assert summary.cell(4, 0).text == 'FINAL REMAINING TOTAL PAYABLE'
    assert [r.cells[1].text for r in summary.rows] == [
        'PHP 5,000.00', 'PHP 1,000.00', 'PHP 0.00', 'PHP 6,000.00', 'PHP 0.00',
    ]
    assert case == before
    assert template().read_bytes() == template_before


@pytest.mark.parametrize('frequency,dates,blocked,final_date', (
    ('daily', ('2027-01-30', '2027-01-31', '2027-02-01'),
     ('2027-02-02', '2027-02-03'), '2027-02-04'),
    ('weekly', ('2027-01-17', '2027-01-24', '2027-01-31'),
     ('2027-02-07', '2027-02-14'), '2027-02-21'),
    ('semi_monthly', ('2027-01-15', '2027-01-30', '2027-02-15'),
     ('2027-02-28', '2027-03-15'), '2027-03-30'),
    ('monthly', ('2027-01-31', '2027-02-28', '2027-03-31'),
     ('2027-04-30', '2027-05-31'), '2027-06-30'),
))
def test_no_collection_tail_skips_only_explicit_blocked_dates(
    frequency, dates, blocked, final_date,
):
    ncs = import_module('gilbic_backend.no_collection_schedule')
    rows = tuple(ncs.OperationalInstallment(
        installment_id=i + 1, installment_number=i + 1,
        contractual_due_date=date.fromisoformat(value),
        effective_due_date=date.fromisoformat(value),
        contractual_amount=Decimal('400.00'), allocated_amount=Decimal('25.00'),
    ) for i, value in enumerate(dates))
    before = deepcopy(rows)
    target = rows[-1].effective_due_date
    excluded = tuple(date.fromisoformat(value) for value in blocked)
    args = dict(installments=rows, no_collection_date=target, payment_frequency=frequency)
    shifts = ncs.plan_no_collection_shift(**args, blocked_dates=excluded)
    # Repeated Sunday/holiday classification cannot create a second adjustment.
    assert shifts == ncs.plan_no_collection_shift(
        **args, blocked_dates=(*excluded, *excluded, target, target),
    )
    assert len(shifts) == 1
    assert shifts[0].new_effective_due_date == date.fromisoformat(final_date)
    assert shifts[0].contractual_due_date == target
    assert shifts[0].prior_effective_due_date == target
    assert shifts[0].contractual_amount == Decimal('400.00')
    assert shifts[0].installment_id == rows[-1].installment_id
    assert rows == before  # No signed date or existing allocation is mutated.


def test_consecutive_no_collection_collision_fails_closed_without_rewriting_rows():
    ncs = import_module('gilbic_backend.no_collection_schedule')
    rows = tuple(ncs.OperationalInstallment(
        installment_id=i + 1, installment_number=i + 1,
        contractual_due_date=date(2027, 1, 30) + timedelta(days=i),
        effective_due_date=date(2027, 1, 30) + timedelta(days=i),
        contractual_amount=Decimal('50.00'),
    ) for i in range(3))
    before = deepcopy(rows)
    with pytest.raises(ncs.NoCollectionScheduleError, match='lose installment order'):
        ncs.plan_no_collection_shift(
            installments=rows, no_collection_date=date(2027, 1, 30),
            payment_frequency='daily',
            blocked_dates=(date(2027, 1, 31), date(2027, 2, 1)),
        )
    assert rows == before
    with pytest.raises(ncs.NoCollectionScheduleError, match='No installment'):
        ncs.plan_no_collection_shift(
            installments=rows, no_collection_date=date(2027, 2, 7),
            payment_frequency='daily',
        )
    assert rows == before


@pytest.mark.parametrize('cash,keep_shift,status', (
    ('30.00', True, 'partial_shifted_prepayment'),
    ('40.00', False, 'full_voluntary_completion'),
))
def test_sunday_holiday_voluntary_exception_keeps_existing_7x7_rule(cash, keep_shift, status):
    # This is the existing 7x7 rule, NOT a new Regular allocation policy.
    nc = import_module('gilbic_backend.seven_by_seven_no_collection_voluntary')
    sunday_holiday = date(2027, 1, 31)  # Synthetic holiday; no statutory calendar claim.
    assert sunday_holiday.weekday() == 6
    affected = nc.NoCollectionAffectedInstallment(
        installment_id=1, installment_number=1,
        contractual_amount=Decimal('50.00'), prepaid_amount=Decimal('10.00'),
    )
    before = deepcopy(affected)
    result = nc.plan_seven_by_seven_no_collection_voluntary_payment(
        transaction_amount=Decimal(cash), collection_date=sunday_holiday,
        no_collection_date=sunday_holiday, past_due_obligations=(),
        affected_installment=affected,
    )
    assert result.status == status
    assert result.keep_no_collection_shift is keep_shift
    assert result.keep_interest_holiday is keep_shift
    assert result.receipt_amount == Decimal(cash)
    assert result.affected_total_after == Decimal(cash) + Decimal('10.00')
    assert sum((r.amount_applied for r in result.instructions), Decimal('0.00')) == Decimal(cash)
    assert affected == before
