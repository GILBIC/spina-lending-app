import hashlib
import io
import json
from datetime import date, timedelta
from pathlib import Path
from decimal import Decimal
from docx import Document
from gilbic_backend import first_loan_documents as documents
from gilbic_backend.first_loan_annex import (
    bind_annex_docx,
    project_packet_schedule,
    ANNEX_R2_SHA256,
)
from uuid import uuid4

import pytest
from pypdf import PdfReader
from reportlab.pdfgen.canvas import Canvas

from gilbic_backend.first_loan_documents import (
    KINDS,
    REQUIRED_FIELDS,
    render_packet,
    template_bundle,
)
from gilbic_backend.first_loan_repository import FirstLoanConflict
from gilbic_backend.first_loan_terms import snapshot_digest


def synthetic_template(extra=()):
    output = io.BytesIO()
    canvas = Canvas(output, pagesize=(595, 842), invariant=1)
    canvas.setFont("Helvetica-Bold", 15)
    canvas.drawString(40, 800, "Synthetic legal template for development testing")
    canvas.setFont("Helvetica", 9)
    for index, name in enumerate(sorted(REQUIRED_FIELDS | set(extra))):
        y = 755 - index * 49
        canvas.drawString(40, y, name.replace("_", " ").title())
        canvas.acroForm.textfield(
            name=name, x=40, y=y - 26, width=510, height=22, fontSize=10
        )
    canvas.showPage()
    canvas.save()
    return output.getvalue()


def packet_fixture(tmp_path, monkeypatch, *, count=12, extra=()):
    template = tmp_path / "approved-synthetic.pdf"
    template.write_bytes(synthetic_template(extra))
    annex = (
        Path(__file__).resolve().parents[2]
        / "tools/annex_a_pdf_proof/assets/SPINA_Schedule_of_Payments_Annex_A_R2_Black_White_Draft.docx"
    )
    if not annex.exists():
        annex = (
            Path(__file__).resolve().parents[3]
            / "counsel-2026-09-19/SPINA_T05B_7x7_Annex_A.docx"
        )
    assert annex.is_file(), "Exact approved Annex A fixture asset is required"
    manifest = {
        "version": "SYNTHETIC-ONLY-1",
        "approved_for_execution": True,
        "annex_a": {"path": str(annex), "sha256": ANNEX_R2_SHA256},
        "agreement_annex_mode": "separate_complete_annex",
        "disclosure_rules": {
            product: {
                "allocation": "SYNTHETIC approved allocation fixture",
                "post_maturity": "SYNTHETIC approved post-maturity fixture",
                "interest_basis": "SYNTHETIC approved interest basis fixture",
            }
            for product in ("regular", "seven_by_seven")
        },
        "documents": {
            kind: {
                "path": str(template),
                "sha256": hashlib.sha256(template.read_bytes()).hexdigest(),
            }
            for kind in KINDS
        },
    }
    path = tmp_path / "templates.json"
    path.write_text(json.dumps(manifest), encoding="utf8")
    monkeypatch.setenv("GILBIC_FIRST_LOAN_TEMPLATE_MANIFEST", str(path))
    _, _, digest = template_bundle()
    packet = {
        "packet_id": str(uuid4()),
        "cif_version_id": str(uuid4()),
        "cif_version_number": 1,
        "borrower": {
            "full_name": "Synthetic Borrower",
            "present_address": "Synthetic address only",
            "phone_number": "00000000000",
            "email": None,
        },
        "product_name": "Synthetic regular product",
        "net_cash": "1000.00",
        "total_deductions": "0.00",
        "privacy": {
            "review_snapshot": {
                "facts": {
                    "registered_lender_name": "Synthetic Lender",
                    "registered_office_address": "Synthetic Office",
                },
                "notice": {"version": "SYN-N"},
                "consent": {"version": "SYN-C"},
            }
        },
        "template": {"version": manifest["version"], "content_sha256": digest},
        "application": {"information": {"request": {"purpose": "Synthetic test only"}}},
        "terms": {
            "product_code": "regular",
            "installment_count": count,
            "first_due_date": "2026-09-20",
            "principal": str(Decimal("80.00") * count),
            "installment_amount": "100.00",
            "payment_frequency": "daily",
            "schedule_basis_date": "2026-09-19",
            "interest_rate_percent": "20.00",
            "daily_interest_per_1000": None,
            "deductions": [],
        },
        "schedule": [
            {
                "installment_number": i + 1,
                "due_date": (date(2026, 9, 20) + timedelta(days=i)).isoformat(),
                "principal_component": "80.00",
                "interest_component": "20.00",
                "contractual_amount": "100.00",
            }
            for i in range(count)
        ],
    }
    return (
        {
            "loan_number": "SYNTHETIC-LOAN-1",
            "approved_by_user_id": str(uuid4()),
            "approved_at": "2026-09-19T00:00:00Z",
            "pricing_snapshot": {},
            "packet": packet,
            "packet_hash": snapshot_digest(packet),
        },
        path,
        template,
    )


def test_packet_fills_three_controlled_documents_and_complete_schedule_without_editable_fields(
    tmp_path, monkeypatch
):
    record, _, _ = packet_fixture(tmp_path, monkeypatch, count=104)

    def synthetic_converter(document):
        doc = Document(io.BytesIO(document))
        output = io.BytesIO()
        canvas = Canvas(output, pagesize=(612, 936))
        y = 900
        for table in doc.tables:
            for row in table.rows:
                line = " | ".join(cell.text for cell in row.cells)
                canvas.drawString(20, y, line[:120])
                y -= 14
                if y < 30:
                    canvas.showPage()
                    y = 900
        canvas.showPage()
        canvas.save()
        return output.getvalue()

    monkeypatch.setattr(documents, "convert_annex_pdf", synthetic_converter)
    monkeypatch.setattr(documents, "verify_arial_pdf", lambda path: None)
    content = render_packet(record)
    pdf = PdfReader(io.BytesIO(content))
    assert len(pdf.pages) > 4
    assert not pdf.get_fields()
    assert all(
        a.get_object().get("/Subtype") != "/Widget"
        for p in pdf.pages
        for a in p.get("/Annots", [])
    )
    text = "\n".join(p.extract_text() for p in pdf.pages)
    assert text.count("Synthetic Borrower") >= 4
    assert record["packet"]["schedule"][-1]["due_date"] in text
    assert "104" in text
    assert record["packet"]["packet_id"] in text


@pytest.mark.parametrize(
    "mutation", ["asset", "version", "packet", "unapproved", "missing_manifest"]
)
def test_issuance_rejects_changed_unapproved_or_unconfigured_material(
    tmp_path, monkeypatch, mutation
):
    record, path, template = packet_fixture(tmp_path, monkeypatch)
    if mutation == "asset":
        template.write_bytes(template.read_bytes() + b"changed")
    elif mutation == "version":
        data = json.loads(path.read_text())
        data["version"] = "changed"
        path.write_text(json.dumps(data))
    elif mutation == "packet":
        record["packet"]["borrower"]["full_name"] = "Another borrower"
    elif mutation == "unapproved":
        data = json.loads(path.read_text())
        data["approved_for_execution"] = False
        path.write_text(json.dumps(data))
    else:
        monkeypatch.delenv("GILBIC_FIRST_LOAN_TEMPLATE_MANIFEST")
    with pytest.raises(FirstLoanConflict):
        render_packet(record)


def test_unsupported_financial_field_is_never_invented(tmp_path, monkeypatch):
    record, _, _ = packet_fixture(
        tmp_path, monkeypatch, extra={"unknown_effective_rate"}
    )
    with pytest.raises(FirstLoanConflict, match="unsupported financial values"):
        render_packet(record)


def test_hash_approved_but_malformed_pdf_template_fails_as_conflict(
    tmp_path, monkeypatch
):
    record, path, template = packet_fixture(tmp_path, monkeypatch)
    template.write_bytes(b"%PDF-1.7\nmalformed controlled template\n%%EOF")
    manifest = json.loads(path.read_text())
    for item in manifest["documents"].values():
        item["sha256"] = hashlib.sha256(template.read_bytes()).hexdigest()
    path.write_text(json.dumps(manifest))
    record["packet"]["template"]["content_sha256"] = template_bundle()[2]
    record["packet_hash"] = snapshot_digest(record["packet"])
    with pytest.raises(FirstLoanConflict, match="PDF"):
        render_packet(record)


def test_malformed_converter_pdf_fails_as_conflict(tmp_path, monkeypatch):
    from gilbic_backend import first_loan_annex as annex

    executable = tmp_path / "synthetic-converter.exe"
    executable.write_bytes(b"test adapter")
    monkeypatch.setenv("GILBIC_OFFICE_DOCUMENT_CONVERTER", str(executable))

    def malformed_conversion(command, **kwargs):
        source = Path(command[-1])
        source.with_suffix(".pdf").write_bytes(b"%PDF-1.7\ninvalid output\n%%EOF")

    monkeypatch.setattr(annex.subprocess, "run", malformed_conversion)
    with pytest.raises(FirstLoanConflict, match="conversion"):
        annex.convert_annex_pdf(b"synthetic converter input")


@pytest.mark.parametrize("count", [1, 7, 8, 104])
def test_annex_binds_every_approved_row_and_remaining_payable_in_retained_layout(
    tmp_path, monkeypatch, count
):
    record, _, _ = packet_fixture(tmp_path, monkeypatch, count=count)
    manifest, files, _ = template_bundle()
    content = bind_annex_docx(
        record,
        documents.packet_fields(record),
        files["annex_a"],
        manifest["disclosure_rules"],
    )
    doc = Document(io.BytesIO(content))
    schedule_tables = [table for table in doc.tables if len(table.columns) == 7]
    rows = [row for table in schedule_tables for row in table.rows[1:]]
    assert len(rows) == count
    assert [row.cells[0].text for row in rows] == [str(i) for i in range(1, count + 1)]
    assert rows[-1].cells[6].text == "0.00"
    assert all(
        table.cell(0, 6).text == "Remaining Total Payable*" for table in schedule_tables
    )
    assert sum(Decimal(row.cells[2].text.replace(",", "")) for row in rows) == Decimal(
        record["packet"]["terms"]["principal"]
    )
    from gilbic_backend.annex_a_template_binding import all_paragraphs

    text = "\n".join(paragraph.text for paragraph in all_paragraphs(doc))
    assert "{" not in text and "}" not in text
    assert record["packet"]["cif_version_id"] in text
    assert "Borrower Signature / Date" in text
    assert "SYNTHETIC approved post-maturity fixture" in text
    assert doc.sections[0].page_width / 914400 == 8
    assert doc.sections[0].page_height / 914400 == 13


def test_template_bundle_locks_legal_rule_text_and_rejects_unapproved_annex(
    tmp_path, monkeypatch
):
    record, path, _ = packet_fixture(tmp_path, monkeypatch)
    old = template_bundle()[2]
    manifest = json.loads(path.read_text())
    manifest["disclosure_rules"]["regular"]["allocation"] += " changed"
    path.write_text(json.dumps(manifest))
    assert template_bundle()[2] != old
    with pytest.raises(FirstLoanConflict, match="do not match"):
        render_packet(record)
    manifest["annex_a"]["sha256"] = "0" * 64
    path.write_text(json.dumps(manifest))
    with pytest.raises(FirstLoanConflict):
        template_bundle()


def test_inconsistent_rows_and_missing_converter_fail_closed(tmp_path, monkeypatch):
    from gilbic_backend.first_loan_annex import convert_annex_pdf

    record, _, _ = packet_fixture(tmp_path, monkeypatch)
    record["packet"]["schedule"][0]["principal_component"] = "0.01"
    with pytest.raises(FirstLoanConflict):
        project_packet_schedule(record["packet"])
    monkeypatch.delenv("GILBIC_OFFICE_DOCUMENT_CONVERTER", raising=False)
    with pytest.raises(FirstLoanConflict, match="converter"):
        convert_annex_pdf(b"not converted")


def test_7x7_annex_requires_and_prints_exact_reviewed_pricing_tuple(
    tmp_path, monkeypatch
):
    record, _, _ = packet_fixture(tmp_path, monkeypatch, count=8)
    record["packet"]["terms"]["product_code"] = "seven_by_seven"
    record["packet"]["product_name"] = "Synthetic 7x7"
    record["packet_hash"] = snapshot_digest(record["packet"])
    manifest, files, _ = template_bundle()
    values = documents.packet_fields(record)
    with pytest.raises(FirstLoanConflict, match="reviewed 7x7 pricing"):
        bind_annex_docx(record, values, files["annex_a"], manifest["disclosure_rules"])
    record["pricing_snapshot"] = {
        "seven_by_seven_penalty_policy": {
            "policy_version": "synthetic-policy",
            "review_id": "synthetic-review",
            "terms_fingerprint": "a" * 64,
            "contractual_monthly_rate": "0.01",
            "proration_days": 30,
            "legal_rate_ceiling": "0.02",
            "lifetime_nonprincipal_cost_ceiling": "640.00",
            "counted_nonprincipal_cost_at_contract_lock": "160.00",
        }
    }
    content = bind_annex_docx(
        record, values, files["annex_a"], manifest["disclosure_rules"]
    )
    doc = Document(io.BytesIO(content))
    from gilbic_backend.annex_a_template_binding import all_paragraphs

    text = "\n".join(paragraph.text for paragraph in all_paragraphs(doc))
    assert "Scheduled Remaining Principal*" in text
    assert "Remaining Total Payable" not in text
    assert "synthetic-review" in text and "proration days: 30" in text
    assert "counted nonprincipal cost at contract lock: 160.00" in text
