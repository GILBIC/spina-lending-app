"""Bind exact approved rows to the retained Annex A R2 layout, without a new engine."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
import hashlib
import io
import os
from pathlib import Path
import subprocess
import tempfile

from docx import Document

from .annex_a_schedule_projection import (
    project_regular_annex_a,
    project_seven_by_seven_annex_a,
)
from .annex_a_template_binding import (
    all_paragraphs,
    replace_tokens,
    set_text,
    fill_table,
    _display_rows,
)
from .contract_schedule_engine import RegularContractInstallment
from .seven_by_seven_signed_schedule import SevenBySevenSignedInstallment
from .document_font_gate import verify_arial_pdf
from .first_loan_repository import FirstLoanConflict

ANNEX_R2_SHA256 = "80ef81aec3141f9c96a5d78f1edd1e7367c8a6be9ab7dca92e81b07756217ddb"


def project_packet_schedule(packet):
    """Validate the immutable approved rows; never regenerate or repair them."""
    try:
        product = packet["terms"]["product_code"]
        if product not in ("regular", "seven_by_seven"):
            raise ValueError("unknown product")
        cls = (
            RegularContractInstallment
            if product == "regular"
            else SevenBySevenSignedInstallment
        )
        source = tuple(
            cls(
                installment_number=row["installment_number"],
                due_date=date.fromisoformat(row["due_date"]),
                contractual_amount=Decimal(row["contractual_amount"]),
                principal_component=Decimal(row["principal_component"]),
                interest_component=Decimal(row["interest_component"]),
            )
            for row in packet["schedule"]
        )
        if not source:
            raise ValueError("empty schedule")
        projector = (
            project_regular_annex_a
            if product == "regular"
            else project_seven_by_seven_annex_a
        )
        return projector(
            original_principal=Decimal(packet["terms"]["principal"]),
            installments=source,
            expected_installment_count=packet["terms"].get("installment_count")
            or len(source),
            expected_first_due_date=date.fromisoformat(
                packet["terms"]["first_due_date"]
            ),
            expected_maturity_date=source[-1].due_date,
            expected_total_payable=sum(
                (row.contractual_amount for row in source), Decimal("0.00")
            ),
        )
    except (ValueError, TypeError, KeyError, ArithmeticError) as error:
        raise FirstLoanConflict(
            "The approved Annex A source rows are incomplete or inconsistent."
        ) from error


def bind_annex_docx(record, values, template: bytes, rules: dict) -> bytes:
    if hashlib.sha256(template).hexdigest() != ANNEX_R2_SHA256:
        raise FirstLoanConflict("The exact approved Annex A R2 layout is required.")
    packet = record["packet"]
    terms = packet["terms"]
    projection = project_packet_schedule(packet)
    regular = terms["product_code"] == "regular"
    product_rules = rules[terms["product_code"]]
    rows = _display_rows(
        {"projection": projection, "product": "Regular" if regular else "7x7"}
    )
    pricing = record.get("pricing_snapshot")
    if not regular and not pricing:
        raise FirstLoanConflict(
            "The exact reviewed 7x7 pricing disclosure is required before issuance."
        )
    penalty = product_rules["post_maturity"]
    if not regular:
        try:
            policy = pricing["seven_by_seven_penalty_policy"]
            policy_keys = (
                "policy_version",
                "review_id",
                "terms_fingerprint",
                "contractual_monthly_rate",
                "proration_days",
                "legal_rate_ceiling",
                "lifetime_nonprincipal_cost_ceiling",
                "counted_nonprincipal_cost_at_contract_lock",
            )
            penalty += "\nReviewed loan-specific policy: " + "; ".join(
                f"{key.replace('_', ' ')}: {policy[key]}" for key in policy_keys
            )
        except (TypeError, KeyError) as error:
            raise FirstLoanConflict(
                "The exact reviewed 7x7 pricing disclosure is incomplete."
            ) from error
    try:
        approver = record["approved_by_user_id"]
        approved_at = record["approved_at"]
        cif_reference = (
            f"{packet['cif_version_id']} / version {packet['cif_version_number']}"
        )
    except KeyError as error:
        raise FirstLoanConflict(
            "The approved Annex A version context is incomplete."
        ) from error
    if hasattr(approved_at, "isoformat"):
        approved_at = approved_at.isoformat()
    doc = Document(io.BytesIO(template))
    tables = doc.tables
    if len(tables) != 6 or [len(table.columns) for table in tables] != [
        2,
        7,
        7,
        2,
        2,
        2,
    ]:
        raise FirstLoanConflict(
            "The approved Annex A structure does not match its binding contract."
        )
    fill_table(tables[1], rows[:7])
    if len(rows) > 7:
        fill_table(tables[2], rows[7:])
    else:
        tables[2]._element.getparent().remove(tables[2]._element)
    for table in (tables[3], tables[5]):
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    rate = (
        f"{terms['interest_rate_percent']}% fixed contractual interest over the approved term"
        if regular
        else f"PHP {projection.rows[0].interest_component:,.2f} per contractual paying day"
    )
    mapping = {
        "{Controlled SEC company identity and registered office address}": f"{values['lender_name']} / {values['lender_address']}",
        "{Borrower Full Name from Locked Loan}": values["borrower_name"],
        "{Borrower Full Name}": values["borrower_name"],
        "{Loan Account No.}": values["loan_number"],
        "{Locked CIF Version}": cif_reference,
        "{Packet ID}": values["packet_reference"],
        "{Packet ID and Version}": f"{values['packet_reference']} / {record['packet_hash']}",
        "{Schedule ID and Version}": f"{values['packet_reference']} / {record['packet_hash']}",
        "{Regular Cash Loan or 7x7 Cash Loan}": values["product_name"],
        "{Approved Gross Principal}": f"{Decimal(terms['principal']):,.2f}",
        "{Exact Approved Rate and Period / Daily Amount}": rate,
        "{Exact Approved Interest Basis and Method}": product_rules["interest_basis"],
        "{Exact Approved Payment Frequency}": values["payment_frequency"].replace(
            "_", " "
        ),
        "{Exact Approved Contractual Term}": f"{len(rows)} installments from {values['first_payment_date']} through {values['maturity_date']}",
        "{Agreed Installment Amount}": f"{Decimal(terms['installment_amount']):,.2f}",
        "{N - Complete Approved Installment Count}": str(len(rows)),
        "{First Contractual Due Date}": values["first_payment_date"],
        "{Last Contractual Installment Date}": values["maturity_date"],
        "{Total under Complete Contractual Schedule}": f"{projection.total_due:,.2f}",
        "{Total Principal Components}": f"{projection.total_principal:,.2f}",
        "{Total Interest Components}": f"{projection.total_interest:,.2f}",
        "{Total Scheduled Charges / 0.00}": "0.00",
        "{Sum of All Contractual Installments}": f"{projection.total_due:,.2f}",
        "{Final Scheduled Remaining Principal}": rows[-1][-1],
        "{Exact Product-Specific Allocation from Locked Disclosure}": product_rules[
            "allocation"
        ],
        "{Exact Applicable Locked Disclosure Rule / None}": penalty,
        "{Authenticated Approver / Locked Approval Reference}": f"{approver} / {values['packet_reference']}",
        "{Document ID}": f"{values['loan_number']}-ANNEX-A",
        "{Template Version}": "Annex A R2",
        "{Packet Version}": record["packet_hash"],
        "{Schedule Version}": record["packet_hash"],
        "{Server Date/Time and Zone}": str(approved_at),
    }
    if terms["payment_frequency"] == "custom":
        mapping = {
            "PHP {Agreed Installment Amount}": "Variable; see complete schedule",
            **mapping,
        }
    for paragraph in list(doc.paragraphs):
        text = paragraph.text
        if text.startswith("Working layout R2") or text.startswith(
            "Template instruction:"
        ):
            paragraph._element.getparent().remove(paragraph._element)
        elif text == "SYSTEM-ALIGNED DRAFT FOR COUNSEL / ACCOUNTING REVIEW":
            set_text(
                paragraph, "CONTRACTUAL SCHEDULE FOR THE IDENTIFIED APPROVED PACKET"
            )
        elif text == "II. CONTRACTUAL INSTALLMENTS - FIRST-PAGE LAYOUT":
            set_text(paragraph, "II. CONTRACTUAL INSTALLMENTS")
        elif text.startswith("All monetary values are in Philippine pesos"):
            set_text(
                paragraph,
                "All monetary values are in Philippine pesos (PHP). Every approved installment is included below and on the continuation pages.",
            )
        elif len(rows) <= 7 and (
            text.startswith("II. CONTRACTUAL INSTALLMENTS - CONTINUATION")
            or text.startswith("Loan: ")
            or paragraph._p.xpath('.//w:br[@w:type="page"]')
        ):
            paragraph._element.getparent().remove(paragraph._element)
        elif text.startswith("Row N is the actual"):
            set_text(paragraph, text.replace("Row N", f"Row {len(rows)}"))
        elif text == "IV. BORROWER ACKNOWLEDGMENT" or text.startswith(
            "I acknowledge receipt and review of the complete Schedule"
        ):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
    for paragraph in all_paragraphs(doc):
        replace_tokens(paragraph, mapping)
        if regular:
            if paragraph.text.startswith("* Scheduled Remaining Principal:"):
                set_text(
                    paragraph,
                    "* Remaining Total Payable: Includes scheduled principal and interest remaining after this installment, assuming all scheduled payments are made fully and on time. This is not a live account balance or payoff quote and is not proof of payment.",
                )
            else:
                replace_tokens(
                    paragraph,
                    {
                        "Scheduled Remaining Principal*": "Remaining Total Payable*",
                        "FINAL SCHEDULED REMAINING PRINCIPAL": "FINAL REMAINING TOTAL PAYABLE",
                    },
                )
    if any(
        "{" in paragraph.text or "}" in paragraph.text
        for paragraph in all_paragraphs(doc)
    ):
        raise FirstLoanConflict("Unresolved Annex A fields prevent issuance.")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def convert_annex_pdf(document: bytes) -> bytes:
    from fitz import FileDataError

    executable = Path(os.getenv("GILBIC_OFFICE_DOCUMENT_CONVERTER", ""))
    if not executable.is_absolute() or not executable.is_file():
        raise FirstLoanConflict("The controlled document converter is not configured.")
    try:
        with tempfile.TemporaryDirectory(prefix="spina-private-annex-") as workspace:
            root = Path(workspace)
            source = root / "annex.docx"
            source.write_bytes(document)
            profile = (root / "profile").as_uri()
            options = {}
            if os.name == "nt":
                startup = subprocess.STARTUPINFO()
                startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
                startup.wShowWindow = subprocess.SW_HIDE
                options["startupinfo"] = startup
                options["creationflags"] = subprocess.CREATE_NO_WINDOW
            subprocess.run(
                [
                    str(executable),
                    f"-env:UserInstallation={profile}",
                    "--headless",
                    "--norestore",
                    "--convert-to",
                    "pdf:writer_pdf_Export",
                    "--outdir",
                    str(root),
                    str(source),
                ],
                check=True,
                timeout=60,
                capture_output=True,
                **options,
            )
            result = root / "annex.pdf"
            verify_arial_pdf(result)
            return result.read_bytes()
    except (OSError, ValueError, FileDataError, subprocess.SubprocessError) as error:
        raise FirstLoanConflict(
            "Annex A conversion or embedded Arial verification failed."
        ) from error
